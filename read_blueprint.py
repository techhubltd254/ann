from docx import Document
import sys

doc = Document(r'C:\Users\lolen\Desktop\kicc\Kenya_National_Exhibition_Platform_Refined_Blueprint_v3.0.docx')

# Write to a text file to avoid encoding issues
with open(r'C:\Users\lolen\Desktop\kicc\blueprint_text.txt', 'w', encoding='utf-8') as f:
    for para in doc.paragraphs:
        f.write(para.text + '\n')
    
    # Also extract tables
    for i, table in enumerate(doc.tables):
        f.write(f'\n=== TABLE {i+1} ===\n')
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            f.write(' | '.join(cells) + '\n')

print("Done. File saved.")
